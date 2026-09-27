"""Testes unitários para o módulo de CH efetiva (sandbox/dae/ch_efetiva.py).

Testa a carga, validações e transformações de C15 sobre a planilha sintética
e variantes mutadas em tmp_path.
"""

from __future__ import annotations

from pathlib import Path
import openpyxl
import pandas as pd
import pytest

from ch_efetiva import (
    ABA_CH_DISCIPLINA,
    COLUNAS_CH_EFETIVA,
    PADRAO_TURMA,
    carregar_ch_efetiva,
    casar_disciplinas,
    ch_da_turma,
    divergencias_calendario,
    nota_sabados,
    normalizar_disciplina,
    resumo_por_carga,
)
from calendario import carregar_calendario


def test_colunas_c15_e_sem_professor(planilha_ch_sintetica: Path) -> None:
    """Verifica se as colunas devolvidas são exatamente as de C15 e sem Professor(a)."""
    df = carregar_ch_efetiva(planilha_ch_sintetica)

    # 1. Colunas devolvidas exatamente as de C15
    assert tuple(df.columns) == COLUNAS_CH_EFETIVA

    # 1b. Colunas derivadas descartadas e disciplina_norm presente
    assert "diferenca" not in df.columns
    assert "pct_nominal" not in df.columns
    assert "disciplina_norm" in df.columns
    assert df.loc[0, "disciplina_norm"] == "TOPOGRAFIA"

    # 2. Sem coluna de professor
    for col in df.columns:
        assert "professor" not in col.lower()

    # 3. Nem o texto do professor fictício em df.to_string()
    str_df = df.to_string().lower()
    assert "fictício" not in str_df
    assert "ficticio" not in str_df
    assert "professor" not in str_df


def test_subgrupo_travessao_para_vazio(planilha_ch_sintetica: Path) -> None:
    """Verifica se o subgrupo '—' (travessão) é normalizado para string vazia ""."""
    df = carregar_ch_efetiva(planilha_ch_sintetica)

    # Turma EST-2A possui '—' na planilha sintética
    linhas_est = df[df["turma"] == "EST-2A"]
    assert len(linhas_est) > 0
    assert (linhas_est["subgrupo"] == "").all()

    # Não deve haver caractere '—' em nenhuma linha de subgrupo
    assert "—" not in df["subgrupo"].values
    assert "-" not in df["subgrupo"].values


def test_serie_e_letra_extraidas(planilha_ch_sintetica: Path) -> None:
    """Verifica se serie e letra de EST/TT-2A são extraídas como 2 e 'A'."""
    df = carregar_ch_efetiva(planilha_ch_sintetica)

    linhas_est_tt = df[df["turma"] == "EST/TT-2A"]
    assert len(linhas_est_tt) > 0
    assert (linhas_est_tt["serie"] == 2).all()
    assert (linhas_est_tt["letra"] == "A").all()

    # Tipagem estrita: serie é int, letra é str
    primeira = linhas_est_tt.iloc[0]
    assert int(primeira["serie"]) == 2
    assert isinstance(primeira["letra"], str)
    assert primeira["letra"] == "A"


def test_ch_origem_planilha_vs_recalculada(
    planilha_ch_sintetica: Path,
    planilha_ch_sintetica_formulas: Path,
) -> None:
    """Verifica attrs['ch_origem'] e igualdade de valores entre versão numérica e com fórmulas."""
    df_num = carregar_ch_efetiva(planilha_ch_sintetica)
    df_form = carregar_ch_efetiva(planilha_ch_sintetica_formulas)

    assert df_num.attrs.get("ch_origem") == "planilha"
    assert df_form.attrs.get("ch_origem") == "recalculada"

    # Ambas devem conter exatamente as mesmas cargas horárias e os mesmos dados
    pd.testing.assert_frame_equal(df_num, df_form, check_like=True)


def test_erro_cabecalho_trocado(tmp_path: Path, planilha_ch_sintetica: Path) -> None:
    """Verifica erro com aba e linha para cabeçalho trocado na linha 1."""
    caminho_mutado = tmp_path / "cabecalho_trocado.xlsx"
    wb = openpyxl.load_workbook(planilha_ch_sintetica)
    ws = wb[ABA_CH_DISCIPLINA]
    # Altera nome de coluna na linha 1
    ws.cell(row=1, column=12, value="Aulas_Semana")
    wb.save(caminho_mutado)

    with pytest.raises(ValueError) as exc:
        carregar_ch_efetiva(caminho_mutado)

    msg = str(exc.value)
    assert ABA_CH_DISCIPLINA in msg
    assert "linha 1" in msg.lower()
    assert "cabeçalho" in msg.lower()


def test_erro_aula_negativa(tmp_path: Path, planilha_ch_sintetica: Path) -> None:
    """Verifica erro com aba e linha para quantidade de aulas negativa (-1)."""
    caminho_mutado = tmp_path / "aula_negativa.xlsx"
    wb = openpyxl.load_workbook(planilha_ch_sintetica)
    ws = wb[ABA_CH_DISCIPLINA]
    # Insere -1 na coluna SEG (col 7) da linha 3
    ws.cell(row=3, column=7, value=-1)
    wb.save(caminho_mutado)

    with pytest.raises(ValueError) as exc:
        carregar_ch_efetiva(caminho_mutado)

    msg = str(exc.value)
    assert ABA_CH_DISCIPLINA in msg
    assert "linha 3" in msg.lower()
    assert "negativa" in msg.lower()


def test_erro_aula_fracionaria(tmp_path: Path, planilha_ch_sintetica: Path) -> None:
    """Verifica erro com aba e linha para aula fracionária (1,5 ou 1.5)."""
    caminho_mutado = tmp_path / "aula_fracionaria.xlsx"
    wb = openpyxl.load_workbook(planilha_ch_sintetica)
    ws = wb[ABA_CH_DISCIPLINA]
    # Insere '1,5' na coluna TER (col 8) da linha 4
    ws.cell(row=4, column=8, value="1,5")
    wb.save(caminho_mutado)

    with pytest.raises(ValueError) as exc:
        carregar_ch_efetiva(caminho_mutado)

    msg = str(exc.value)
    assert ABA_CH_DISCIPLINA in msg
    assert "linha 4" in msg.lower()
    assert "inteira" in msg.lower()


def test_erro_aulas_sem_inconsistente(tmp_path: Path, planilha_ch_sintetica: Path) -> None:
    """Verifica erro com aba e linha para 'Aulas/sem' inconsistente com a soma dos dias."""
    caminho_mutado = tmp_path / "aulas_sem_inconsistente.xlsx"
    wb = openpyxl.load_workbook(planilha_ch_sintetica)
    ws = wb[ABA_CH_DISCIPLINA]
    # Altera 'Aulas/sem' (col 12) da linha 2 de 2 para 4
    ws.cell(row=2, column=12, value=4)
    wb.save(caminho_mutado)

    with pytest.raises(ValueError) as exc:
        carregar_ch_efetiva(caminho_mutado)

    msg = str(exc.value)
    assert ABA_CH_DISCIPLINA in msg
    assert "linha 2" in msg.lower()
    assert "aulas/sem" in msg.lower()


def test_erro_ch_nominal_inconsistente(
    tmp_path: Path,
    planilha_ch_sintetica: Path,
) -> None:
    """Verifica erro com aba e linha para 'CH nominal' ≠ aulas/sem × 40."""
    caminho_mutado = tmp_path / "ch_nominal_inconsistente.xlsx"
    wb = openpyxl.load_workbook(planilha_ch_sintetica)
    ws = wb[ABA_CH_DISCIPLINA]
    # Linha 2 tem aulas/sem = 2 (CH nominal 80). Mudamos CH nominal (col 13) para 81.
    ws.cell(row=2, column=13, value=81)
    wb.save(caminho_mutado)

    with pytest.raises(ValueError) as exc:
        carregar_ch_efetiva(caminho_mutado)

    msg = str(exc.value)
    assert ABA_CH_DISCIPLINA in msg
    assert "linha 2" in msg.lower()
    assert "ch nominal" in msg.lower()


def test_erro_ch_efetiva_inconsistente_com_bimestres(
    tmp_path: Path,
    planilha_ch_sintetica: Path,
) -> None:
    """Verifica erro com aba e linha para 'CH efetiva' ≠ soma dos bimestres."""
    caminho_mutado = tmp_path / "ch_efetiva_inconsistente.xlsx"
    wb = openpyxl.load_workbook(planilha_ch_sintetica)
    ws = wb[ABA_CH_DISCIPLINA]
    # Linha 2 soma bimestres 20+20+18+18 = 76. Mudamos CH efetiva (col 18) para 80.
    ws.cell(row=2, column=18, value=80)
    wb.save(caminho_mutado)

    with pytest.raises(ValueError) as exc:
        carregar_ch_efetiva(caminho_mutado)

    msg = str(exc.value)
    assert ABA_CH_DISCIPLINA in msg
    assert "linha 2" in msg.lower()
    assert "ch efetiva" in msg.lower()


def test_erro_chave_duplicada(tmp_path: Path, planilha_ch_sintetica: Path) -> None:
    """Verifica erro com aba e linha para oferta duplicada (turma, subgrupo, disciplina)."""
    caminho_mutado = tmp_path / "chave_duplicada.xlsx"
    wb = openpyxl.load_workbook(planilha_ch_sintetica)
    ws = wb[ABA_CH_DISCIPLINA]
    # Clona a linha 2 (EST-2A, —, TOPOGRAFIA) na linha 9
    linha_2_vals = [ws.cell(row=2, column=c).value for c in range(1, 21)]
    ws.append(linha_2_vals)
    wb.save(caminho_mutado)

    with pytest.raises(ValueError) as exc:
        carregar_ch_efetiva(caminho_mutado)

    msg = str(exc.value)
    assert ABA_CH_DISCIPLINA in msg
    assert "linha" in msg.lower()
    assert "chave duplicada" in msg.lower()


def test_erro_turma_fora_padrao(tmp_path: Path, planilha_ch_sintetica: Path) -> None:
    """Verifica erro com aba e linha para turma fora do padrão (ex.: 'EST2A' sem hífen)."""
    caminho_mutado = tmp_path / "turma_invalida.xlsx"
    wb = openpyxl.load_workbook(planilha_ch_sintetica)
    ws = wb[ABA_CH_DISCIPLINA]
    # Altera turma da linha 2 para EST2A
    ws.cell(row=2, column=2, value="EST2A")
    wb.save(caminho_mutado)

    with pytest.raises(ValueError) as exc:
        carregar_ch_efetiva(caminho_mutado)

    msg = str(exc.value)
    assert ABA_CH_DISCIPLINA in msg
    assert "linha 2" in msg.lower()
    assert "est2a" in msg.lower() or "turma" in msg.lower()


def test_normalizar_disciplina() -> None:
    """Verifica regras de normalização de nomes de disciplinas."""
    # Prefixos de língua estrangeira e sufixo de série
    assert (
        normalizar_disciplina("LÍNGUA ESTRANGEIRA: INGLÊS - 2ª SÉRIE")
        == "INGLES"
    )

    # Sufixo de série e acentos
    assert normalizar_disciplina("FÍSICA - 2ª SÉRIE") == "FISICA"

    # Espaços múltiplos e maiúsculas
    assert (
        normalizar_disciplina("Máquinas  e Equipamentos")
        == "MAQUINAS E EQUIPAMENTOS"
    )

    # Preservação de distinção de laboratório
    assert normalizar_disciplina("LABORATÓRIO DE SOLOS") != "SOLOS"
    assert normalizar_disciplina("LABORATÓRIO DE SOLOS") == "LABORATORIO DE SOLOS"
    assert normalizar_disciplina("SOLOS") == "SOLOS"


def test_resumo_por_carga_fixture(planilha_ch_sintetica: Path) -> None:
    """Verifica resumo_por_carga calculado sobre a fixture sintética."""
    df = carregar_ch_efetiva(planilha_ch_sintetica)
    resumo = resumo_por_carga(df)

    # A fixture possui 7 linhas: 4 com 70 h/a (<90%), 1 com 74 h/a (90%-95%), 2 com 76 ou 38 h/a (>=95%)
    assert resumo == {"< 90%": 4, "90–95%": 1, "≥ 95%": 2, "total": 7}


def test_divergencias_calendario_fixture_coerente(planilha_ch_sintetica: Path) -> None:
    """Verifica que a fixture sintética coerente não possui divergências com o calendário oficial."""
    cal = carregar_calendario()
    divs = divergencias_calendario(planilha_ch_sintetica, cal)
    assert divs == []


def test_divergencias_calendario_aba_calendario_mutada(
    tmp_path: Path,
    planilha_ch_sintetica: Path,
) -> None:
    """Cópia com 1º BI/QUI = 11 na aba Calendário gera 1 divergência citando Calendário, 1º BI, QUI, 11 × 10."""
    caminho_mutado = tmp_path / "cal_mutado.xlsx"
    wb = openpyxl.load_workbook(planilha_ch_sintetica)
    ws = wb["Calendário"]
    # Na aba Calendário, linha 4 é '1º BI' e coluna 5 é 'QUI'
    ws.cell(row=4, column=5, value=11)
    wb.save(caminho_mutado)

    cal = carregar_calendario()
    divs = divergencias_calendario(caminho_mutado, cal)
    assert len(divs) == 1
    d = divs[0]
    assert "Calendário" in d
    assert "1º BI" in d
    assert "QUI" in d
    assert "11 × 10" in d


def test_divergencias_calendario_linha_bimestre_alterado(
    tmp_path: Path,
    planilha_ch_sintetica: Path,
) -> None:
    """Linha com 2º BI alterado de 18 para 19 gera divergência citando turma, disciplina, 2º BI, 19 × 18 e sem nome do professor."""
    caminho_mutado = tmp_path / "linha_mutada.xlsx"
    wb = openpyxl.load_workbook(planilha_ch_sintetica)
    ws = wb[ABA_CH_DISCIPLINA]
    # Linha 8: TT-2A, OPERAÇÃO DE TRANSPORTES, 2º BI = 18 (coluna 15).
    # Ajustamos também CH efetiva (coluna 18) de 74 para 75 para manter integridade da soma
    ws.cell(row=8, column=15, value=19)
    ws.cell(row=8, column=18, value=75)
    wb.save(caminho_mutado)

    cal = carregar_calendario()
    divs = divergencias_calendario(caminho_mutado, cal)
    assert len(divs) == 1
    d = divs[0]
    assert "TT-2A" in d
    assert "OPERAÇÃO DE TRANSPORTES" in d
    assert "2º BI" in d
    assert "19 × 18" in d
    # Não pode conter o nome do professor fictício
    assert "fictício" not in d.lower()
    assert "ficticio" not in d.lower()
    assert "professor" not in d.lower()


def test_divergencias_calendario_faixa_invalida(planilha_ch_sintetica: Path) -> None:
    """Linha de 2 aulas com CH efetiva 80 (fora de 70-76) gera divergência de faixa."""
    df = carregar_ch_efetiva(planilha_ch_sintetica)
    df_mut = df.copy()
    # Linha 0 (EST-2A, TOPOGRAFIA) tem aulas_sem = 2. Alteramos ch_efetiva de 76 para 80
    df_mut.loc[0, "ch_efetiva"] = 80

    cal = carregar_calendario()
    divs = divergencias_calendario(df_mut, cal)
    assert len(divs) == 1
    d = divs[0]
    assert "EST-2A" in d
    assert "TOPOGRAFIA" in d
    assert "faixa" in d.lower()
    assert "80" in d
    assert "70–76" in d or "70-76" in d


def test_nota_sabados() -> None:
    """Verifica conteúdo das notas de sábados letivos para Estradas e Hospedagem."""
    cal = carregar_calendario()
    nota_est = nota_sabados(cal, "Estradas")
    assert "23/05" in nota_est
    assert "2º" in nota_est
    assert "cenário A" in nota_est

    nota_hosp = nota_sabados(cal, "Hospedagem")
    assert "26/09" in nota_hosp


def test_ch_da_turma(planilha_ch_sintetica: Path) -> None:
    """Verifica filtragem de ofertas de CH por curso, série e turma (C17)."""
    df = carregar_ch_efetiva(planilha_ch_sintetica)

    # 1. Estradas série 2 pega EST-2A + EST/TT-2A e não TT-2A
    res_est = ch_da_turma(df, "Estradas", 2, "TÉCNICO EM ESTRADAS - BH-1EST - A (2025)")
    turmas_est = set(res_est["turma"])
    assert turmas_est == {"EST-2A", "EST/TT-2A"}
    assert "TT-2A" not in turmas_est

    # 2. "Trânsito" pega TT-2A + EST/TT-2A e não EST-2A
    res_tt = ch_da_turma(df, "Trânsito", 2, "TÉCNICO EM TRÂNSITO - BH-1TRANS - A (2025)")
    turmas_tt = set(res_tt["turma"])
    assert turmas_tt == {"TT-2A", "EST/TT-2A"}
    assert "EST-2A" not in turmas_tt

    # 3. Série 3 → vazio
    res_s3 = ch_da_turma(df, "Estradas", 3, "TÉCNICO EM ESTRADAS - BH-1EST - A (2025)")
    assert len(res_s3) == 0


def test_casar_disciplinas_legenda_ficticia(planilha_ch_sintetica: Path) -> None:
    """Verifica casamento com legenda fictícia, subgrupos congruentes e divergentes (C17)."""
    df = carregar_ch_efetiva(planilha_ch_sintetica)
    df_turma = ch_da_turma(df, "Estradas", 2, "TÉCNICO EM ESTRADAS - BH-1EST - A (2025)")

    legenda_ficticia = {
        "C1": "LÍNGUA ESTRANGEIRA: INGLÊS - 2ª SÉRIE",
        "C2": "REDAÇÃO - 2ª SÉRIE",
        "C3": "EDUCAÇÃO FÍSICA - 2ª SÉRIE",
    }

    casadas, sem_linha = casar_disciplinas(legenda_ficticia, df_turma)

    # C1 e C2 casados com ch_bim_1..4 / arranjo corretos
    assert "C1" in casadas
    assert "C2" in casadas

    c1 = casadas["C1"]
    assert c1["ch_bim_1"] == 18
    assert c1["ch_bim_2"] == 18
    assert c1["ch_bim_3"] == 18
    assert c1["ch_bim_4"] == 16
    assert c1["arranjo"] == {"SEX": 2}

    # C2 (T1 = T2) sem subgrupos_divergentes
    c2 = casadas["C2"]
    assert c2["ch_bim_1"] == 20
    assert c2["ch_bim_2"] == 20
    assert c2["ch_bim_3"] == 16
    assert c2["ch_bim_4"] == 14
    assert c2["arranjo"] == {"SEG": 2}
    assert c2["subgrupos_divergentes"] is False

    # C3 na lista de sem-linha (nomes da legenda sem linha na planilha)
    assert "EDUCAÇÃO FÍSICA - 2ª SÉRIE" in sem_linha

    # Variante com T1 ≠ T2 → a menor CH e subgrupos_divergentes True
    df_mut = df_turma.copy()
    idx_t2 = df_mut[df_mut["subgrupo"] == "T2"].index[0]
    df_mut.loc[idx_t2, "seg"] = 0
    df_mut.loc[idx_t2, "sex"] = 2
    df_mut.loc[idx_t2, "ch_bim_1"] = 18
    df_mut.loc[idx_t2, "ch_efetiva"] = 68

    casadas_mut, _ = casar_disciplinas(legenda_ficticia, df_mut)
    c2_mut = casadas_mut["C2"]
    assert c2_mut["subgrupos_divergentes"] is True
    assert c2_mut["ch_efetiva"] == 68
    assert c2_mut["ch_bim_1"] == 18
    assert c2_mut["arranjo"] == {"SEX": 2}

